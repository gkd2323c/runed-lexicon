#!/usr/bin/env python3
"""Export a human-readable review sheet (xlsx) for a Skyrim MOD translation.

The translated xTranslator XML holds every string but no dialogue structure.
The Mutagen dialogue context holds the structure but no translation. This tool
joins the two and writes one sheet a person can actually read: dialogue grouped
by quest and ordered by its place in the plugin, then non-dialogue rows grouped
by record field. Untranslated cells are highlighted.

Rows are matched to the dialogue context in two passes. Most records carry a
real EDID; INFO records without one carry a bracketed FormID placeholder
instead (`[06006C3A]`), which is matched against the context's FormID. The
second pass typically doubles dialogue coverage, so keep both.

Every project input is read-only; the only file written is the workbook.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
from collections import Counter
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["序", "段", "任务", "类别", "对话", "说话人", "记录", "字段", "原文", "译文", "状态"]
WIDTHS = [6, 8, 30, 14, 34, 18, 32, 14, 60, 60, 8]
WRAP_COLUMNS = {4, 8, 9}  # 对话 / 原文 / 译文 需要换行
UNTRANSLATED_FILL = PatternFill("solid", start_color="D9D2E9")
HEADER_FILL = PatternFill("solid", start_color="3D3D3D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALIGN_WRAP = Alignment(vertical="top", wrap_text=True)
ALIGN_PLAIN = Alignment(vertical="top")

FORMID_EDID = re.compile(r"^\[([0-9A-Fa-f]{8})\]$")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--root", default=".", help="项目根目录（默认当前目录）")
    parser.add_argument("--mod", required=True, help="mods/ 下的 MOD 目录名，例如 TheKalpicAnomaly_VIGILANT.esp")
    parser.add_argument("--work", default=None, help=".work/ 下的工作目录名（默认取 MOD 名去掉扩展名）")
    parser.add_argument("--xml", default=None, help="显式指定译后 XML 路径")
    parser.add_argument("--context", default=None, help="显式指定 dialogue-context JSON 路径")
    parser.add_argument("--out", default=None, help="输出 xlsx 路径")
    parser.add_argument("--dialogue-only", action="store_true", help="只导出对话段，丢弃非对话行")
    return parser.parse_args(argv)


def load_xml_rows(xml_path: Path) -> list[dict]:
    """Read every <String> in document order, resolving bracketed FormID EDIDs."""
    rows: list[dict] = []
    for _, element in ET.iterparse(str(xml_path), events=("end",)):
        if element.tag != "String":
            continue
        edid = (element.findtext("EDID") or "").strip()
        match = FORMID_EDID.match(edid)
        rows.append(
            {
                "edid": edid,
                "formid": match.group(1).upper() if match else "",
                "rec": (element.findtext("REC") or "").strip(),
                "source": element.findtext("Source") or "",
                "dest": element.findtext("Dest") or "",
                "doc": len(rows),
            }
        )
        element.clear()
    return rows


def format_speakers(candidates) -> str:
    """Flatten speaker_candidates into a display string; tolerate shape drift."""
    names: list[str] = []
    for candidate in candidates or []:
        if isinstance(candidate, str):
            if candidate:
                names.append(candidate)
        elif isinstance(candidate, dict):
            for key in ("name", "editor_id", "display_name", "npc_name", "npc", "form_id"):
                value = candidate.get(key)
                if value:
                    names.append(str(value))
                    break
    return "; ".join(dict.fromkeys(names))


def build_indexes(context: dict) -> dict:
    """Index the dialogue context by EDID and by FormID."""
    dialogues: dict[str, dict] = {}
    infos: dict[str, dict] = {}
    dialogues_by_formid: dict[str, dict] = {}
    infos_by_formid: dict[str, dict] = {}
    quest_rank: dict[str, int] = {}

    for seq, dialogue in enumerate(context.get("dialogues") or []):
        quest = (dialogue.get("quest_edid") or "").strip()
        if quest and quest not in quest_rank:
            quest_rank[quest] = len(quest_rank)

        kind = " / ".join(
            part
            for part in ((dialogue.get("category") or "").strip(), (dialogue.get("subtype") or "").strip())
            if part
        )
        edid = (dialogue.get("editor_id") or "").strip()
        entry = {"edid": edid, "quest": quest, "kind": kind, "seq": seq}
        if edid:
            dialogues[edid] = entry
        form_id = (dialogue.get("form_id") or "").strip().upper()
        if form_id:
            dialogues_by_formid.setdefault(form_id, entry)

        for info_seq, info in enumerate(dialogue.get("infos") or []):
            info_edid = (info.get("editor_id") or "").strip()
            info_formid = (info.get("form_id") or "").strip().upper()
            info_entry = {
                "dialogue": edid,
                "seq": info_seq,
                "speaker": format_speakers(info.get("speaker_candidates")),
            }
            if info_edid and info_edid not in infos:
                infos[info_edid] = info_entry
            if info_formid and info_formid not in infos_by_formid:
                infos_by_formid[info_formid] = info_entry

    return {
        "dialogues": dialogues,
        "infos": infos,
        "dialogues_by_formid": dialogues_by_formid,
        "infos_by_formid": infos_by_formid,
        "quest_rank": quest_rank,
    }


def locate(row: dict, indexes: dict) -> tuple[str, dict | None, str]:
    """Resolve a row to the dialogue context. Returns (kind, entry, how)."""
    edid = row["edid"]
    formid = row["formid"]
    if edid in indexes["infos"]:
        return "info", indexes["infos"][edid], "edid"
    if edid in indexes["dialogues"]:
        return "dialogue", indexes["dialogues"][edid], "edid"
    if formid and formid in indexes["infos_by_formid"]:
        return "info", indexes["infos_by_formid"][formid], "formid"
    if formid and formid in indexes["dialogues_by_formid"]:
        return "dialogue", indexes["dialogues_by_formid"][formid], "formid"
    return "", None, "unmatched"


def pick_xml(mod_dir: Path, base: str, explicit: str | None) -> Path:
    if explicit:
        return Path(explicit)
    for candidate in (mod_dir / f"{base}_english_chinese_translated.xml", mod_dir / f"{base}_english_chinese.xml"):
        if candidate.is_file():
            return candidate
    fallback = sorted(mod_dir.glob("*translated*.xml")) or sorted(mod_dir.glob("*.xml"))
    if fallback:
        return fallback[0]
    raise FileNotFoundError(f"在 {mod_dir} 下找不到可用的 XML")


def pick_context(root: Path, work_name: str, explicit: str | None) -> Path:
    if explicit:
        return Path(explicit)
    context_dir = root / ".work" / work_name / "context"
    if context_dir.is_dir():
        found = sorted(context_dir.glob("*-dialogue-context.json"))
        if found:
            return found[0]
    available = sorted(path.parent.parent.name for path in (root / ".work").glob("*/context/*-dialogue-context.json"))
    hint = "、".join(available) if available else "（无）"
    raise FileNotFoundError(f"在 .work/{work_name}/context/ 下找不到 dialogue-context；现有工作目录：{hint}")


def main(argv: list[str] | None = None) -> int:
    started = time.perf_counter()
    args = parse_args(argv)
    root = Path(args.root).resolve()

    mod_dir = root / "mods" / args.mod
    if not mod_dir.is_dir():
        print(f"[error] MOD 目录不存在：{mod_dir}", file=sys.stderr)
        return 2

    base = args.mod
    for extension in (".esp", ".esm", ".esl"):
        if base.lower().endswith(extension):
            base = base[: -len(extension)]
            break
    work_name = args.work or base

    try:
        xml_path = pick_xml(mod_dir, base, args.xml)
        context_path = pick_context(root, work_name, args.context)
    except FileNotFoundError as error:
        print(f"[error] {error}", file=sys.stderr)
        return 2

    context = json.loads(Path(context_path).read_text(encoding="utf-8"))
    xml_rows = load_xml_rows(xml_path)
    indexes = build_indexes(context)
    quest_titles = {
        row["edid"]: row["dest"]
        for row in xml_rows
        if row["rec"] == "QUST:FULL" and row["edid"] and row["dest"]
    }
    dialogue_titles = {
        row["edid"]: row["dest"]
        for row in xml_rows
        if row["rec"] == "DIAL:FULL" and row["edid"] and row["dest"]
    }

    dialogue_rows: list[dict] = []
    other_rows: list[dict] = []
    counters: Counter = Counter()

    for row in xml_rows:
        kind, entry, how = locate(row, indexes)
        counters[f"{how}_{kind}" if kind else "unmatched"] += 1

        if kind == "info" and entry:
            dialogue_key = entry["dialogue"]
            dialogue = indexes["dialogues"].get(dialogue_key) or {}
            quest = dialogue.get("quest", "")
            dialogue_rows.append(
                {
                    "segment": "对话",
                    "quest": quest_titles.get(quest, quest),
                    "kind": dialogue.get("kind", ""),
                    "dialogue": dialogue_titles.get(dialogue_key, dialogue_key),
                    "speaker": entry["speaker"],
                    "row": row,
                    "sort": (
                        indexes["quest_rank"].get(quest, 10**6),
                        dialogue.get("seq", 10**6),
                        entry["seq"],
                        row["doc"],
                    ),
                }
            )
        elif kind == "dialogue" and entry:
            dialogue_key = entry.get("edid", "") or row["edid"]
            quest = entry.get("quest", "")
            dialogue_rows.append(
                {
                    "segment": "对话",
                    "quest": quest_titles.get(quest, quest),
                    "kind": entry.get("kind", ""),
                    "dialogue": dialogue_titles.get(dialogue_key, dialogue_key),
                    "speaker": "",
                    "row": row,
                    "sort": (indexes["quest_rank"].get(quest, 10**6), entry.get("seq", 10**6), -1, row["doc"]),
                }
            )
        else:
            other_rows.append(
                {
                    "segment": "非对话",
                    "quest": "",
                    "kind": "",
                    "dialogue": "",
                    "speaker": "",
                    "row": row,
                    "sort": (row["rec"], row["edid"], row["doc"]),
                }
            )

    dialogue_rows.sort(key=lambda item: item["sort"])
    other_rows.sort(key=lambda item: item["sort"])
    ordered = dialogue_rows if args.dialogue_only else dialogue_rows + other_rows

    out_path = Path(args.out) if args.out else mod_dir / f"{base}-review-sheet.xlsx"

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("审校稿")
    sheet.freeze_panes = "A2"
    if ordered:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(ordered) + 1}"

    header_cells = []
    for title in HEADERS:
        cell = WriteOnlyCell(sheet, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_PLAIN
        header_cells.append(cell)
    sheet.append(header_cells)
    for position, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width

    untranslated = 0
    for index, item in enumerate(ordered, start=1):
        row = item["row"]
        is_untranslated = row["source"] == row["dest"]
        if is_untranslated:
            untranslated += 1
        values = [
            index,
            item["segment"],
            item["quest"],
            item["kind"],
            item["dialogue"],
            item["speaker"],
            row["edid"],
            row["rec"],
            row["source"],
            row["dest"],
            "未译" if is_untranslated else "",
        ]
        cells = []
        for column, value in enumerate(values):
            cell = WriteOnlyCell(sheet, value=value)
            cell.alignment = ALIGN_WRAP if column in WRAP_COLUMNS else ALIGN_PLAIN
            cells.append(cell)
        if is_untranslated:
            cells[9].fill = UNTRANSLATED_FILL
        sheet.append(cells)

    workbook.save(str(out_path))

    report = {
        "output": str(out_path),
        "xml": str(xml_path),
        "context": str(context_path),
        "xml_strings": len(xml_rows),
        "dialogue_rows": len(dialogue_rows),
        "non_dialogue_rows": len(other_rows),
        "untranslated": untranslated,
        "matched_edid_info": counters["edid_info"],
        "matched_edid_dialogue": counters["edid_dialogue"],
        "matched_formid_info": counters["formid_info"],
        "matched_formid_dialogue": counters["formid_dialogue"],
        "unmatched": counters["unmatched"],
        "index_dialogue_edid": len(indexes["dialogues"]),
        "index_dialogue_formid": len(indexes["dialogues_by_formid"]),
        "index_info_edid": len(indexes["infos"]),
        "index_info_formid": len(indexes["infos_by_formid"]),
        "elapsed_s": round(time.perf_counter() - started, 2),
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
